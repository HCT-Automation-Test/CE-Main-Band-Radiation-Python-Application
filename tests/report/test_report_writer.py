from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.models.enums import (
    AntennaPol,
    DUTAntPosition,
    DUTAxis,
    Modulation,
    NRBand,
    RBAllocation,
    TestChannel as ChannelEnum,
    TestType as MeasurementTestType,
)
from app.models.measurement_result import (
    BilogResult,
    FundamentalResult,
    HarmonicResult,
    InbandResult,
)
from app.models.test_case import TestCase
from app.report.excel_writer import ExcelReportWriter
from app.report.report_writer import ReportWriter
from app.report.screenshot_store import (
    build_fundamental_screenshot_path,
    build_sa_screenshot_path,
)


class _SignalRecorder:
    def __init__(self) -> None:
        self.messages: list[str] = []

    def emit(self, message: str) -> None:
        self.messages.append(message)


class _Events:
    def __init__(self) -> None:
        self.screenshot_saved = _SignalRecorder()


def _make_case(
    *,
    test_type: MeasurementTestType,
    band: NRBand = NRBand.N78,
    bandwidth: float = 100.0,
    antenna_pol: AntennaPol = AntennaPol.V,
    dut_axis: DUTAxis = DUTAxis.X,
    dut_ant_position: DUTAntPosition = DUTAntPosition.UPPER,
) -> TestCase:
    return TestCase(
        band=band,
        bandwidth=bandwidth,
        test_type=test_type,
        modulation=Modulation.QPSK,
        rb_allocation=RBAllocation.INNER_FULL,
        test_channel=ChannelEnum.MID,
        antenna_pol=antenna_pol,
        dut_axis=dut_axis,
        dut_ant_position=dut_ant_position,
    )


def _make_workbook(path: Path, *bands: NRBand) -> Path:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for band in bands:
        workbook.create_sheet(f"Sub6 {band.value}")
    workbook.save(path)
    return path


def _make_writer(
    *,
    report_path: Path,
    screenshot_root: Path,
    events: _Events | None = None,
) -> ReportWriter:
    return ReportWriter(
        excel_writer=ExcelReportWriter(report_path),
        screenshot_root_dir=screenshot_root,
        model_name="SM-000",
        events=events,
    )


def test_write_fundamental_writes_excel_and_stores_fundamental_screenshot(
    tmp_path: Path,
) -> None:
    report_path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    screenshot_root = tmp_path / "screenshots"
    source = tmp_path / "temp" / "fundamental.png"
    source.parent.mkdir()
    source.write_bytes(b"fundamental")
    events = _Events()
    case = _make_case(test_type=MeasurementTestType.HARMONIC)
    writer = _make_writer(
        report_path=report_path,
        screenshot_root=screenshot_root,
        events=events,
    )

    assert writer.write_fundamental(
        case,
        FundamentalResult(
            screenshot_path=source,
            fundamental_frequency_mhz=3550.25,
        ),
    )

    workbook = load_workbook(report_path)
    assert workbook["Sub6 n78"]["J49"].value == 3550.25
    destination = build_fundamental_screenshot_path(
        root_dir=screenshot_root,
        model_name="SM-000",
        test_case=case,
    )
    assert destination.read_bytes() == b"fundamental"
    assert events.screenshot_saved.messages == [str(destination)]


def test_read_fundamental_frequency_delegates_to_excel_writer(tmp_path: Path) -> None:
    report_path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = _make_writer(
        report_path=report_path,
        screenshot_root=tmp_path / "screenshots",
    )
    writer.excel_writer.workbook["Sub6 n78"]["J49"] = 3550.25

    assert (
        writer.read_fundamental_frequency(
            _make_case(test_type=MeasurementTestType.HARMONIC)
        )
        == 3550.25
    )


@pytest.mark.parametrize(
    ("test_type", "result", "expected_cell", "expected_value", "source_name"),
    [
        (
            MeasurementTestType.HARMONIC,
            HarmonicResult(
                harmonic_2nd_peak_level_dbm=-41.0,
                harmonic_3rd_peak_level_dbm=-40.0,
                harmonic_4th_peak_level_dbm=-39.0,
                harmonic_5th_peak_level_dbm=-38.0,
            ),
            "B66",
            -41.0,
            "harmonic.png",
        ),
        (
            MeasurementTestType.INBAND,
            InbandResult(inband_peak_level_dbm=-45.5),
            "B61",
            -45.5,
            "inband.png",
        ),
        (
            MeasurementTestType.BILOG,
            BilogResult(bilog_peak_level_dbm=-38.25),
            "B56",
            -38.25,
            "bilog.png",
        ),
    ],
)
def test_write_measurement_writes_excel_and_stores_sa_screenshot(
    tmp_path: Path,
    test_type: MeasurementTestType,
    result: HarmonicResult | InbandResult | BilogResult,
    expected_cell: str,
    expected_value: float,
    source_name: str,
) -> None:
    report_path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    screenshot_root = tmp_path / "screenshots"
    source = tmp_path / "temp" / source_name
    source.parent.mkdir()
    source.write_bytes(source_name.encode("utf-8"))
    result.screenshot_path = source
    events = _Events()
    case = _make_case(test_type=test_type)
    writer = _make_writer(
        report_path=report_path,
        screenshot_root=screenshot_root,
        events=events,
    )

    if test_type is MeasurementTestType.HARMONIC:
        written = writer.write_harmonic(case, result)
    elif test_type is MeasurementTestType.INBAND:
        written = writer.write_inband(case, result)
    else:
        written = writer.write_bilog(case, result)

    assert written
    workbook = load_workbook(report_path)
    assert workbook["Sub6 n78"][expected_cell].value == expected_value
    destination = build_sa_screenshot_path(
        root_dir=screenshot_root,
        model_name="SM-000",
        test_case=case,
    )
    assert destination.read_bytes() == source_name.encode("utf-8")
    assert events.screenshot_saved.messages == [str(destination)]


def test_write_fails_when_screenshot_path_is_missing(tmp_path: Path) -> None:
    report_path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = _make_writer(
        report_path=report_path,
        screenshot_root=tmp_path / "screenshots",
    )

    with pytest.raises(ValueError, match="screenshot_path is required"):
        writer.write_inband(
            _make_case(test_type=MeasurementTestType.INBAND),
            InbandResult(inband_peak_level_dbm=-45.5),
        )


def test_write_fails_when_screenshot_source_is_missing(tmp_path: Path) -> None:
    report_path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = _make_writer(
        report_path=report_path,
        screenshot_root=tmp_path / "screenshots",
    )

    with pytest.raises(FileNotFoundError):
        writer.write_inband(
            _make_case(test_type=MeasurementTestType.INBAND),
            InbandResult(
                screenshot_path=tmp_path / "temp" / "missing.png",
                inband_peak_level_dbm=-45.5,
            ),
        )


def test_excel_failure_does_not_store_screenshot(tmp_path: Path) -> None:
    report_path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    screenshot_root = tmp_path / "screenshots"
    source = tmp_path / "temp" / "fundamental.png"
    source.parent.mkdir()
    source.write_bytes(b"fundamental")
    writer = _make_writer(
        report_path=report_path,
        screenshot_root=screenshot_root,
    )

    with pytest.raises(ValueError, match="fundamental_frequency_mhz"):
        writer.write_fundamental(
            _make_case(test_type=MeasurementTestType.HARMONIC),
            FundamentalResult(screenshot_path=source),
        )

    assert not screenshot_root.exists()
