from __future__ import annotations

import sys
import types
from dataclasses import replace
from pathlib import Path
from shutil import copy2
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.models.enums import (
    AntennaPol,
    DUTAntPosition,
    DUTAxis,
    Generation,
    LTEBand,
    LTERBAllocation,
    Modulation,
    NRBand,
    RBAllocation,
    TestChannel as ChannelEnum,
    TestType as MeasurementTestType,
)
from app.models.measurement_result import (
    BilogResult,
    FundamentalResult,
    HarmonicResult,
    InbandResult,
)
from app.models.test_case import TestCase
import app.report.excel_writer as excel_writer
from app.report.excel_writer import (
    ExcelReportWriter,
    _find_open_excel_workbook,
    _find_workbook_from_running_object_table,
)
from app.report import template_map


def _make_case(
    *,
    band: NRBand = NRBand.N78,
    bandwidth: float = 100.0,
    test_type: MeasurementTestType = MeasurementTestType.HARMONIC,
    antenna_pol: AntennaPol = AntennaPol.V,
    dut_axis: DUTAxis = DUTAxis.X,
    dut_ant_position: DUTAntPosition = DUTAntPosition.UPPER,
    generation: Generation = Generation.NR_SA,
    test_channel: ChannelEnum = ChannelEnum.MID,
    anchor_band: LTEBand = LTEBand.B1,
) -> TestCase:
    return TestCase(
        band=band,
        bandwidth=bandwidth,
        test_type=test_type,
        modulation=Modulation.QPSK,
        rb_allocation=RBAllocation.INNER_FULL,
        test_channel=test_channel,
        antenna_pol=antenna_pol,
        dut_axis=dut_axis,
        dut_ant_position=dut_ant_position,
        generation=generation,
        anchor_band=anchor_band if generation is Generation.NR_NSA else None,
    )


def _make_lte_case(
    *,
    band: LTEBand = LTEBand.B1,
    bandwidth: float = 20.0,
    test_type: MeasurementTestType = MeasurementTestType.HARMONIC,
    antenna_pol: AntennaPol = AntennaPol.V,
    dut_axis: DUTAxis = DUTAxis.X,
    dut_ant_position: DUTAntPosition = DUTAntPosition.UPPER,
    test_channel: ChannelEnum = ChannelEnum.MID,
    rb_allocation: LTERBAllocation = LTERBAllocation.MID,
) -> TestCase:
    return TestCase(
        band=band,
        bandwidth=bandwidth,
        test_type=test_type,
        modulation=Modulation.QPSK,
        rb_allocation=rb_allocation,
        test_channel=test_channel,
        antenna_pol=antenna_pol,
        dut_axis=dut_axis,
        dut_ant_position=dut_ant_position,
        generation=Generation.LTE,
    )


_LTE_LB_BANDS = frozenset({LTEBand.B5, LTEBand.B8, LTEBand.B20, LTEBand.B28})


def _make_workbook(path: Path, *bands: NRBand) -> Path:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for band in bands:
        workbook.create_sheet(f"Sub6 {band.value}")
    workbook.save(path)
    return path


def _make_lte_workbook(
    path: Path,
    *bands: LTEBand,
    extra_sheets: tuple[str, ...] = (),
) -> Path:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for band in bands:
        worksheet = workbook.create_sheet(f"LTE_BAND {band.value.removeprefix('B')}")
        worksheet["A7"] = band.value
        worksheet["G7"] = 15
        harmonic_rows = (19, 21, 23, 25) if band in _LTE_LB_BANDS else (24, 26, 28, 30)
        for order, row in zip((2, 3, 4, 5), harmonic_rows, strict=True):
            worksheet[f"A{row}"] = f"=I7*{order}"
    for sheet_name in extra_sheets:
        workbook.create_sheet(sheet_name)
    workbook.save(path)
    return path


def _cell(path: Path, sheet: str, cell: str) -> object:
    workbook = load_workbook(path)
    return workbook[sheet][cell].value


class _FakeWorkbook:
    def __init__(self, full_name: Path | str | None = None) -> None:
        if full_name is not None:
            self.FullName = str(full_name)
        self.saved = False
        self.activated = False
        self.Application = _FakeExcelApplication()
        self.sheets: dict[str, _FakeSheet] = {}

    def Worksheets(self, sheet_name: str) -> "_FakeSheet":
        return self.sheets.setdefault(sheet_name, _FakeSheet(sheet_name))

    def Activate(self) -> None:
        self.activated = True

    def Save(self) -> None:
        self.saved = True


class _FakeSheet:
    def __init__(self, name: str = "") -> None:
        self.name = name
        self.activated = False
        self.ranges: dict[str, _FakeRange] = {}

    def Activate(self) -> None:
        self.activated = True

    def Range(self, cell: str) -> "_FakeRange":
        return self.ranges.setdefault(cell, _FakeRange())


class _FakeRange:
    def __init__(self) -> None:
        self.Value: object = None
        self.selected = False
        self.activated = False

    def Select(self) -> None:
        self.selected = True

    def Activate(self) -> None:
        self.activated = True


class _FakeExcelApplication:
    def __init__(self) -> None:
        self.Visible = False


class _BrokenWorkbook:
    @property
    def FullName(self) -> str:
        raise RuntimeError("broken workbook")


class _FakeRunningObjects:
    def __init__(self, objects: list[Any]) -> None:
        self._objects = objects

    def EnumRunning(self) -> "_FakeMonikers":
        return _FakeMonikers(list(range(len(self._objects))))

    def GetObject(self, moniker: int) -> Any:
        return self._objects[moniker]


class _FakeMonikers:
    def __init__(self, monikers: list[int]) -> None:
        self._monikers = monikers
        self._index = 0

    def Next(self, _: int) -> tuple[int, ...]:
        if self._index >= len(self._monikers):
            return ()
        moniker = self._monikers[self._index]
        self._index += 1
        return (moniker,)


class _FakePythoncom:
    def __init__(self, objects: list[Any]) -> None:
        self._running_objects = _FakeRunningObjects(objects)

    def CreateBindCtx(self, _: int) -> object:
        return object()

    def GetRunningObjectTable(self) -> _FakeRunningObjects:
        return self._running_objects


class _FakeWin32Client:
    def __init__(
        self,
        active_excel: Any | None = None,
        dispatch_excel: Any | None = None,
    ) -> None:
        self._active_excel = active_excel
        self._dispatch_excel = dispatch_excel

    def Dispatch(self, candidate: Any) -> Any:
        if candidate == "Excel.Application" and self._dispatch_excel is not None:
            return self._dispatch_excel
        return candidate

    def GetActiveObject(self, _: str) -> Any:
        if self._active_excel is None:
            raise RuntimeError("no active excel")
        return self._active_excel


class _FakeWorkbooks(list[Any]):
    def __init__(self, workbooks: list[Any]) -> None:
        super().__init__(workbooks)
        self.open_calls: list[str] = []

    def Open(self, path: str) -> Any:
        self.open_calls.append(path)
        workbook = _FakeWorkbook(path)
        self.append(workbook)
        return workbook


class _FakeExcel:
    def __init__(self, workbooks: list[Any]) -> None:
        self.Visible = False
        self.Workbooks = _FakeWorkbooks(workbooks)


def _install_fake_com_modules(
    monkeypatch: pytest.MonkeyPatch,
) -> tuple[types.ModuleType, types.ModuleType]:
    pythoncom = types.ModuleType("pythoncom")
    com_calls: list[str] = []

    def co_initialize() -> None:
        com_calls.append("initialize")

    def co_uninitialize() -> None:
        com_calls.append("uninitialize")

    pythoncom.CoInitialize = co_initialize  # type: ignore[attr-defined]
    pythoncom.CoUninitialize = co_uninitialize  # type: ignore[attr-defined]
    pythoncom.calls = com_calls  # type: ignore[attr-defined]
    win32com = types.ModuleType("win32com")
    win32_client = types.ModuleType("win32com.client")
    win32com.client = win32_client  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "pythoncom", pythoncom)
    monkeypatch.setitem(sys.modules, "win32com", win32com)
    monkeypatch.setitem(sys.modules, "win32com.client", win32_client)
    return pythoncom, win32_client


def test_write_fundamental_saves_default_cells_to_template_file(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    case = _make_case()

    assert writer.write_fundamental(
        case,
        FundamentalResult(fundamental_frequency_mhz=3550.25),
    )

    sheet = "Sub6 n78"
    workbook = load_workbook(path)
    worksheet = workbook[sheet]
    assert worksheet["A49"].value == "n78"
    assert worksheet["B49"].value is None
    assert worksheet["C49"].value == 636666
    assert worksheet["D49"].value == 100
    assert worksheet["E49"].value == case.rb_size
    assert worksheet["F49"].value == case.rb_offset
    assert worksheet["G49"].value == "QPSK"
    assert worksheet["H49"].value == 30
    assert worksheet["I49"].value == "Upper"
    assert worksheet["J49"].value == 3550.25
    assert worksheet["A66"].value == 7100.5
    assert worksheet["A68"].value == 10650.75
    assert worksheet["A70"].value == 14201
    assert worksheet["A72"].value == 17751.25


def test_write_fundamental_writes_resolved_low_and_high_channels(
    tmp_path: Path,
) -> None:
    low_path = _make_workbook(tmp_path / "low.xlsx", NRBand.N78)
    low_writer = ExcelReportWriter(low_path)
    low_writer.write_fundamental(
        _make_case(test_channel=ChannelEnum.LOW),
        FundamentalResult(fundamental_frequency_mhz=3350.01),
    )
    assert _cell(low_path, "Sub6 n78", "C49") == 623334

    high_path = _make_workbook(tmp_path / "high.xlsx", NRBand.N78)
    high_writer = ExcelReportWriter(high_path)
    high_writer.write_fundamental(
        _make_case(test_channel=ChannelEnum.HIGH),
        FundamentalResult(fundamental_frequency_mhz=3750.0),
    )
    assert _cell(high_path, "Sub6 n78", "C49") == 650000


def test_write_fundamental_raises_for_missing_channel_lookup(
    tmp_path: Path,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N28)
    writer = ExcelReportWriter(path)
    case = _make_case(
        band=NRBand.N28,
        bandwidth=25.0,
        test_channel=ChannelEnum.MID,
    )

    with pytest.raises(ValueError, match="No NR channel lookup"):
        writer.write_fundamental(
            case,
            FundamentalResult(fundamental_frequency_mhz=725.5),
        )


def test_read_fundamental_frequency_returns_none_for_empty_cell(
    tmp_path: Path,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)

    assert writer.read_fundamental_frequency(_make_case()) is None


def test_read_fundamental_frequency_reads_default_layout_cell(
    tmp_path: Path,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    writer.workbook["Sub6 n78"]["J49"] = 3550.25

    assert writer.read_fundamental_frequency(_make_case()) == 3550.25


def test_read_fundamental_frequency_reads_lb_layout_cell(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N5)
    writer = ExcelReportWriter(path)
    writer.workbook["Sub6 n5"]["J44"] = 836

    assert (
        writer.read_fundamental_frequency(_make_case(band=NRBand.N5, bandwidth=20.0))
        == 836.0
    )


@pytest.mark.parametrize("value", [True, "3550.25", "=SUM(1, 2)"])
def test_read_fundamental_frequency_rejects_invalid_cell_values(
    tmp_path: Path,
    value: object,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    writer.workbook["Sub6 n78"]["J49"] = value

    with pytest.raises(ValueError, match="Invalid fundamental frequency value"):
        writer.read_fundamental_frequency(_make_case())


def test_write_harmonic_saves_default_axis_and_pol_cells(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    case = _make_case(antenna_pol=AntennaPol.H, dut_axis=DUTAxis.Y)

    assert writer.write_harmonic(
        case,
        HarmonicResult(
            harmonic_2nd_peak_level_dbm=-41.0,
            harmonic_3rd_peak_level_dbm=-39.5,
            harmonic_4th_peak_level_dbm=-38.0,
            harmonic_5th_peak_level_dbm=-37.5,
        ),
    )

    sheet = "Sub6 n78"
    assert _cell(path, sheet, "C67") == -41.0
    assert _cell(path, sheet, "C69") == -39.5
    assert _cell(path, sheet, "C71") == -38.0
    assert _cell(path, sheet, "C73") == -37.5


def test_write_fundamental_uses_common_harmonic_frequency_cells_for_pol(
    tmp_path: Path,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    case = _make_case(antenna_pol=AntennaPol.H)

    writer.write_fundamental(
        case,
        FundamentalResult(fundamental_frequency_mhz=3550.25),
    )

    workbook = load_workbook(path)
    worksheet = workbook["Sub6 n78"]
    assert worksheet["A66"].value == 7100.5
    assert worksheet["A67"].value is None
    assert worksheet["A68"].value == 10650.75
    assert worksheet["A69"].value is None


def test_write_harmonic_writes_dash_for_missing_levels(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    case = _make_case(antenna_pol=AntennaPol.H, dut_axis=DUTAxis.Y)

    assert writer.write_harmonic(
        case,
        HarmonicResult(
            harmonic_2nd_peak_level_dbm=None,
            harmonic_3rd_peak_level_dbm=-39.5,
            harmonic_4th_peak_level_dbm=None,
            harmonic_5th_peak_level_dbm=-37.5,
        ),
    )

    sheet = "Sub6 n78"
    assert _cell(path, sheet, "C67") == "-"
    assert _cell(path, sheet, "C69") == -39.5
    assert _cell(path, sheet, "C71") == "-"
    assert _cell(path, sheet, "C73") == -37.5


def test_write_fundamental_and_harmonic_use_lb_layout(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N5)
    writer = ExcelReportWriter(path)
    case = _make_case(
        band=NRBand.N5,
        bandwidth=20.0,
        dut_ant_position=DUTAntPosition.LOWER,
    )

    writer.write_fundamental(case, FundamentalResult(fundamental_frequency_mhz=836.5))
    writer.write_harmonic(
        case,
        HarmonicResult(
            harmonic_2nd_peak_level_dbm=-42.0,
            harmonic_3rd_peak_level_dbm=-41.0,
            harmonic_4th_peak_level_dbm=-40.0,
            harmonic_5th_peak_level_dbm=-39.0,
        ),
    )

    workbook = load_workbook(path)
    worksheet = workbook["Sub6 n5"]
    assert worksheet["A44"].value == "n5"
    assert worksheet["B44"].value is None
    assert worksheet["C44"].value == 167300
    assert worksheet["D44"].value == 20
    assert worksheet["H44"].value == 15
    assert worksheet["I44"].value == "Lower"
    assert worksheet["J44"].value == 836.5
    assert worksheet["A56"].value == 1673
    assert worksheet["A58"].value == 2509.5
    assert worksheet["B56"].value == -42.0
    assert worksheet["B62"].value == -39.0


def test_write_inband_and_bilog_default_cells(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)

    inband_case = _make_case(
        test_type=MeasurementTestType.INBAND,
        antenna_pol=AntennaPol.H,
        dut_axis=DUTAxis.Z,
    )
    bilog_case = _make_case(
        test_type=MeasurementTestType.BILOG,
        antenna_pol=AntennaPol.V,
        dut_axis=DUTAxis.X,
    )

    writer.write_inband(inband_case, InbandResult(inband_peak_level_dbm=-45.5))
    writer.write_bilog(bilog_case, BilogResult(bilog_peak_level_dbm=-38.25))

    sheet = "Sub6 n78"
    assert _cell(path, sheet, "D62") == -45.5
    assert _cell(path, sheet, "B56") == -38.25


def test_write_inband_uses_lb_layout(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N5)
    writer = ExcelReportWriter(path)
    case = _make_case(
        band=NRBand.N5,
        bandwidth=20.0,
        test_type=MeasurementTestType.INBAND,
        antenna_pol=AntennaPol.H,
        dut_axis=DUTAxis.Z,
    )

    writer.write_inband(case, InbandResult(inband_peak_level_dbm=-47.0))

    assert _cell(path, "Sub6 n5", "D52") == -47.0


def test_write_nsa_default_cells_use_shifted_rows(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    base_case = _make_case(generation=Generation.NR_NSA)
    inband_case = _make_case(
        test_type=MeasurementTestType.INBAND,
        antenna_pol=AntennaPol.H,
        dut_axis=DUTAxis.Z,
        generation=Generation.NR_NSA,
    )
    bilog_case = _make_case(
        test_type=MeasurementTestType.BILOG,
        generation=Generation.NR_NSA,
    )

    writer.write_fundamental(
        base_case,
        FundamentalResult(fundamental_frequency_mhz=3550.25),
    )
    writer.write_harmonic(
        base_case,
        HarmonicResult(
            harmonic_2nd_peak_level_dbm=-41.0,
            harmonic_3rd_peak_level_dbm=-40.0,
            harmonic_4th_peak_level_dbm=-39.0,
            harmonic_5th_peak_level_dbm=-38.0,
        ),
    )
    writer.write_inband(inband_case, InbandResult(inband_peak_level_dbm=-45.5))
    writer.write_bilog(bilog_case, BilogResult(bilog_peak_level_dbm=-38.25))

    worksheet = load_workbook(path)["Sub6 n78"]
    assert worksheet["A7"].value == "n78"
    assert worksheet["B7"].value == "B1"
    assert worksheet["C7"].value == 636666
    assert worksheet["H7"].value == 30
    assert worksheet["J7"].value == 3550.25
    assert worksheet["A24"].value == 7100.5
    assert worksheet["B24"].value == -41.0
    assert worksheet["D20"].value == -45.5
    assert worksheet["B14"].value == -38.25


def test_write_nsa_lb_cells_use_shifted_rows(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N5)
    writer = ExcelReportWriter(path)
    case = _make_case(
        band=NRBand.N5,
        bandwidth=20.0,
        generation=Generation.NR_NSA,
    )
    inband_case = _make_case(
        band=NRBand.N5,
        bandwidth=20.0,
        test_type=MeasurementTestType.INBAND,
        antenna_pol=AntennaPol.H,
        dut_axis=DUTAxis.Z,
        generation=Generation.NR_NSA,
    )

    writer.write_fundamental(case, FundamentalResult(fundamental_frequency_mhz=836.5))
    writer.write_harmonic(
        case,
        HarmonicResult(
            harmonic_2nd_peak_level_dbm=-42.0,
            harmonic_3rd_peak_level_dbm=-41.0,
            harmonic_4th_peak_level_dbm=-40.0,
            harmonic_5th_peak_level_dbm=-39.0,
        ),
    )
    writer.write_inband(inband_case, InbandResult(inband_peak_level_dbm=-47.0))

    worksheet = load_workbook(path)["Sub6 n5"]
    assert worksheet["A7"].value == "n5"
    assert worksheet["B7"].value == "B1"
    assert worksheet["C7"].value == 167300
    assert worksheet["H7"].value == 15
    assert worksheet["J7"].value == 836.5
    assert worksheet["A19"].value == 1673
    assert worksheet["B19"].value == -42.0
    assert worksheet["D15"].value == -47.0


def test_band_info_is_written_once_per_band(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    first_case = _make_case(test_type=MeasurementTestType.INBAND)
    second_case = _make_case(
        bandwidth=90.0,
        test_type=MeasurementTestType.INBAND,
        antenna_pol=AntennaPol.H,
    )

    writer.write_inband(first_case, InbandResult(inband_peak_level_dbm=-45.0))
    writer.write_inband(second_case, InbandResult(inband_peak_level_dbm=-44.0))

    assert _cell(path, "Sub6 n78", "D49") == 100


def test_non_fundamental_write_writes_band_and_scs(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)

    writer.write_inband(
        _make_case(test_type=MeasurementTestType.INBAND),
        InbandResult(inband_peak_level_dbm=-45.0),
    )

    worksheet = load_workbook(path)["Sub6 n78"]
    assert worksheet["A49"].value == "n78"
    assert worksheet["B49"].value is None
    assert worksheet["H49"].value == 30


def test_write_fails_when_sheet_is_missing(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N77)
    writer = ExcelReportWriter(path)

    with pytest.raises(ValueError, match="Workbook has no sheet"):
        writer.write_inband(
            _make_case(test_type=MeasurementTestType.INBAND),
            InbandResult(inband_peak_level_dbm=-45.0),
        )


def test_write_fails_for_invalid_cell_coordinate(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    layout = template_map.LAYOUTS["default"]
    monkeypatch.setitem(
        template_map.LAYOUTS,
        "default",
        replace(layout, fundamental_frequency_cell="A1:B2"),
    )

    with pytest.raises(ValueError, match="Invalid single cell coordinate"):
        writer.write_fundamental(
            _make_case(),
            FundamentalResult(fundamental_frequency_mhz=3550.25),
        )


def test_write_inband_and_bilog_write_dash_for_missing_peak(
    tmp_path: Path,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)

    inband_case = _make_case(
        test_type=MeasurementTestType.INBAND,
        antenna_pol=AntennaPol.H,
        dut_axis=DUTAxis.Z,
    )
    bilog_case = _make_case(
        test_type=MeasurementTestType.BILOG,
        antenna_pol=AntennaPol.V,
        dut_axis=DUTAxis.X,
    )

    assert writer.write_inband(
        inband_case,
        InbandResult(inband_peak_level_dbm=None),
    )
    assert writer.write_bilog(bilog_case, BilogResult(bilog_peak_level_dbm=None))

    assert _cell(path, "Sub6 n78", "D62") == "-"
    assert _cell(path, "Sub6 n78", "B56") == "-"


def test_band_info_is_written_once_per_band_and_generation(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    sa_case = _make_case(test_type=MeasurementTestType.INBAND)
    nsa_case = _make_case(
        bandwidth=90.0,
        test_type=MeasurementTestType.INBAND,
        generation=Generation.NR_NSA,
    )

    writer.write_inband(sa_case, InbandResult(inband_peak_level_dbm=-45.0))
    writer.write_inband(nsa_case, InbandResult(inband_peak_level_dbm=-44.0))

    assert _cell(path, "Sub6 n78", "D49") == 100
    assert _cell(path, "Sub6 n78", "D7") == 90


def test_nsa_band_info_keeps_first_anchor_for_writer_lifetime(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    first_case = _make_case(
        test_type=MeasurementTestType.INBAND,
        generation=Generation.NR_NSA,
        anchor_band=LTEBand.B1,
    )
    second_case = _make_case(
        bandwidth=90.0,
        test_type=MeasurementTestType.INBAND,
        antenna_pol=AntennaPol.H,
        generation=Generation.NR_NSA,
        anchor_band=LTEBand.B3,
    )

    writer.write_inband(first_case, InbandResult(inband_peak_level_dbm=-45.0))
    writer.write_inband(second_case, InbandResult(inband_peak_level_dbm=-44.0))

    worksheet = load_workbook(path)["Sub6 n78"]
    assert worksheet["B7"].value == "B1"
    assert worksheet["D7"].value == 100


def test_new_writer_can_record_new_first_nsa_anchor(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    first_writer = ExcelReportWriter(path)
    first_writer.write_inband(
        _make_case(
            test_type=MeasurementTestType.INBAND,
            generation=Generation.NR_NSA,
            anchor_band=LTEBand.B1,
        ),
        InbandResult(inband_peak_level_dbm=-45.0),
    )

    second_writer = ExcelReportWriter(path)
    second_writer.write_inband(
        _make_case(
            bandwidth=90.0,
            test_type=MeasurementTestType.INBAND,
            generation=Generation.NR_NSA,
            anchor_band=LTEBand.B3,
        ),
        InbandResult(inband_peak_level_dbm=-44.0),
    )

    worksheet = load_workbook(path)["Sub6 n78"]
    assert worksheet["B7"].value == "B3"
    assert worksheet["D7"].value == 90


@pytest.mark.parametrize(
    ("anchor_band", "error"),
    [
        (None, "NR_NSA requires anchor_band"),
        (LTEBand.B34, "Unsupported NSA anchor_band"),
    ],
)
def test_nsa_anchor_validation_fails_before_writing_band_info(
    tmp_path: Path,
    anchor_band: LTEBand | None,
    error: str,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    case = _make_case(
        test_type=MeasurementTestType.INBAND,
        generation=Generation.NR_NSA,
    )
    case.anchor_band = anchor_band

    with pytest.raises(ValueError, match=error):
        writer.write_inband(case, InbandResult(inband_peak_level_dbm=-45.0))

    worksheet = writer.workbook["Sub6 n78"]
    assert worksheet["A7"].value is None
    assert worksheet["B7"].value is None


def test_nsa_anchor_mapping_is_required_before_writing_band_info(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    layout = template_map.LAYOUTS["default"]
    monkeypatch.setitem(
        template_map.LAYOUTS,
        "default",
        replace(
            layout,
            band_info_cells=replace(layout.band_info_cells, anchor_band=None),
        ),
    )

    with pytest.raises(ValueError, match="has no anchor_band cell"):
        writer.write_inband(
            _make_case(
                test_type=MeasurementTestType.INBAND,
                generation=Generation.NR_NSA,
            ),
            InbandResult(inband_peak_level_dbm=-45.0),
        )

    worksheet = writer.workbook["Sub6 n78"]
    assert worksheet["A7"].value is None
    assert worksheet["B7"].value is None


def test_write_nsa_anchor_to_actual_template_copy_preserves_sa_band_info(
    tmp_path: Path,
) -> None:
    source = Path(__file__).resolve().parents[2] / "resources/templates/NR_Rawdata_.xlsx"
    path = tmp_path / source.name
    copy2(source, path)
    source_worksheet = load_workbook(source)["Sub6 n7"]
    expected_sa_values = [
        source_worksheet.cell(row=49, column=column).value for column in range(1, 10)
    ]
    writer = ExcelReportWriter(path)

    writer.write_inband(
        _make_case(
            band=NRBand.N7,
            bandwidth=20.0,
            test_type=MeasurementTestType.INBAND,
            generation=Generation.NR_NSA,
            anchor_band=LTEBand.B1,
        ),
        InbandResult(inband_peak_level_dbm=-45.0),
    )

    worksheet = load_workbook(path)["Sub6 n7"]
    actual_sa_values = [
        worksheet.cell(row=49, column=column).value for column in range(1, 10)
    ]
    assert worksheet["A7"].value == "n7"
    assert worksheet["B7"].value == "B1"
    assert actual_sa_values == expected_sa_values


def test_write_falls_back_to_open_excel_workbook_when_file_is_locked(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    fallback_cells: list[tuple[str, str, object]] = []

    def raise_permission_error(_: Path) -> None:
        raise PermissionError("locked")

    def save_open_workbook() -> None:
        fallback_cells.extend(writer._pending_cells)
        writer._pending_cells.clear()

    monkeypatch.setattr(writer.workbook, "save", raise_permission_error)
    monkeypatch.setattr(writer, "_save_open_workbook_with_excel", save_open_workbook)

    assert writer.write_inband(
        _make_case(
            test_type=MeasurementTestType.INBAND,
            antenna_pol=AntennaPol.H,
            dut_axis=DUTAxis.Z,
        ),
        InbandResult(inband_peak_level_dbm=-45.5),
    )

    assert ("Sub6 n78", "D62", -45.5) in fallback_cells
    assert writer._pending_cells == []


def test_prepare_live_monitoring_opens_excel_and_activates_start_sheet(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    pythoncom, win32_client = _install_fake_com_modules(monkeypatch)
    excel = _FakeExcel([])
    fake_client = _FakeWin32Client(dispatch_excel=excel)
    win32_client.GetActiveObject = fake_client.GetActiveObject  # type: ignore[attr-defined]
    win32_client.Dispatch = fake_client.Dispatch  # type: ignore[attr-defined]

    writer.prepare_live_monitoring(_make_case())

    opened_workbook = excel.Workbooks[-1]
    assert excel.Visible
    assert opened_workbook.activated
    assert opened_workbook.Worksheets("Sub6 n78").activated
    assert writer._live_monitoring_enabled
    assert pythoncom.calls == [  # type: ignore[attr-defined]
        "initialize",
        "uninitialize",
    ]


def test_prepare_live_monitoring_reuses_already_open_workbook(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    workbook = _FakeWorkbook(path)
    _install_fake_com_modules(monkeypatch)
    monkeypatch.setattr(
        excel_writer,
        "_find_open_excel_workbook",
        lambda *_: workbook,
    )

    writer.prepare_live_monitoring(_make_case())

    assert workbook.activated
    assert workbook.Worksheets("Sub6 n78").activated
    assert writer._live_monitoring_enabled


def test_live_monitoring_write_uses_excel_and_selects_last_cell(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    workbook = _FakeWorkbook(path)
    writer._live_monitoring_enabled = True
    monkeypatch.setattr(
        excel_writer,
        "_get_or_open_excel_workbook",
        lambda *_: workbook,
    )

    writer.write_inband(
        _make_case(
            test_type=MeasurementTestType.INBAND,
            antenna_pol=AntennaPol.H,
            dut_axis=DUTAxis.Z,
        ),
        InbandResult(inband_peak_level_dbm=-45.5),
    )

    sheet = workbook.Worksheets("Sub6 n78")
    assert sheet.activated
    assert sheet.Range("D62").Value == -45.5
    assert sheet.Range("D62").selected
    assert workbook.saved
    assert writer._pending_cells == []


def test_live_monitoring_switches_sheet_when_band_changes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78, NRBand.N5)
    writer = ExcelReportWriter(path)
    workbook = _FakeWorkbook(path)
    writer._live_monitoring_enabled = True
    monkeypatch.setattr(
        excel_writer,
        "_get_or_open_excel_workbook",
        lambda *_: workbook,
    )

    writer.write_inband(
        _make_case(
            test_type=MeasurementTestType.INBAND,
            antenna_pol=AntennaPol.H,
            dut_axis=DUTAxis.Z,
        ),
        InbandResult(inband_peak_level_dbm=-45.5),
    )
    writer.write_inband(
        _make_case(
            band=NRBand.N5,
            bandwidth=20.0,
            test_type=MeasurementTestType.INBAND,
            antenna_pol=AntennaPol.H,
            dut_axis=DUTAxis.Z,
        ),
        InbandResult(inband_peak_level_dbm=-47.0),
    )

    assert workbook.Worksheets("Sub6 n78").activated
    assert workbook.Worksheets("Sub6 n5").activated
    assert workbook.Worksheets("Sub6 n78").Range("D62").Value == -45.5
    assert workbook.Worksheets("Sub6 n5").Range("D52").Value == -47.0
    assert workbook.Worksheets("Sub6 n5").Range("D52").selected


def test_prepare_live_monitoring_raises_when_excel_cannot_open(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    _install_fake_com_modules(monkeypatch)
    monkeypatch.setattr(excel_writer, "_get_excel_application", lambda *_: object())

    with pytest.raises(RuntimeError, match="live monitoring"):
        writer.prepare_live_monitoring(_make_case())


def test_save_open_workbook_with_excel_writes_cells_to_found_workbook(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    workbook = _FakeWorkbook(path)
    pythoncom, win32_client = _install_fake_com_modules(monkeypatch)

    def find_workbook(
        workbook_path: Path,
        pythoncom_module: object,
        win32_client_module: object,
    ) -> object:
        assert workbook_path == path
        assert pythoncom_module is pythoncom
        assert win32_client_module is win32_client
        return workbook

    monkeypatch.setattr(excel_writer, "_find_open_excel_workbook", find_workbook)
    writer._pending_cells = [("Sub6 n78", "D62", -45.5)]

    writer._save_open_workbook_with_excel()

    assert pythoncom.calls == [  # type: ignore[attr-defined]
        "initialize",
        "uninitialize",
    ]
    assert workbook.Worksheets("Sub6 n78").Range("D62").Value == -45.5
    assert workbook.saved
    assert writer._pending_cells == []


def test_save_open_workbook_with_excel_raises_when_no_workbook_is_found(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _make_workbook(tmp_path / "report.xlsx", NRBand.N78)
    writer = ExcelReportWriter(path)
    pythoncom, _ = _install_fake_com_modules(monkeypatch)
    monkeypatch.setattr(
        excel_writer,
        "_find_open_excel_workbook",
        lambda *_: None,
    )
    writer._pending_cells = [("Sub6 n78", "D62", -45.5)]

    with pytest.raises(PermissionError, match="no open Excel workbook"):
        writer._save_open_workbook_with_excel()

    assert pythoncom.calls == [  # type: ignore[attr-defined]
        "initialize",
        "uninitialize",
    ]


def test_find_open_workbook_prefers_running_object_table(tmp_path: Path) -> None:
    target_path = tmp_path / "report.xlsx"
    rot_workbook = _FakeWorkbook(target_path)
    active_workbook = _FakeWorkbook(tmp_path / "other.xlsx")

    found = _find_open_excel_workbook(
        target_path,
        _FakePythoncom([rot_workbook]),
        _FakeWin32Client(_FakeExcel([active_workbook])),
    )

    assert found is rot_workbook


def test_find_open_workbook_uses_active_excel_when_rot_has_no_match(
    tmp_path: Path,
) -> None:
    target_path = tmp_path / "report.xlsx"
    active_workbook = _FakeWorkbook(target_path)

    found = _find_open_excel_workbook(
        target_path,
        _FakePythoncom([_FakeWorkbook(tmp_path / "other.xlsx")]),
        _FakeWin32Client(_FakeExcel([active_workbook])),
    )

    assert found is active_workbook


def test_running_object_table_ignores_objects_without_readable_full_name(
    tmp_path: Path,
) -> None:
    target_path = tmp_path / "report.xlsx"
    target_workbook = _FakeWorkbook(target_path)

    found = _find_workbook_from_running_object_table(
        target_path,
        _FakePythoncom([_BrokenWorkbook(), target_workbook]),
        _FakeWin32Client(),
    )

    assert found is target_workbook


def test_find_open_workbook_returns_none_when_rot_and_active_excel_do_not_match(
    tmp_path: Path,
) -> None:
    target_path = tmp_path / "report.xlsx"

    found = _find_open_excel_workbook(
        target_path,
        _FakePythoncom([_FakeWorkbook(tmp_path / "other.xlsx")]),
        _FakeWin32Client(_FakeExcel([_FakeWorkbook(tmp_path / "also_other.xlsx")])),
    )

    assert found is None


def test_write_lte_mhb_results_preserves_template_values_and_formulas(
    tmp_path: Path,
) -> None:
    path = _make_lte_workbook(tmp_path / "lte.xlsx", LTEBand.B1)
    writer = ExcelReportWriter(path)
    inband_case = _make_lte_case(
        test_type=MeasurementTestType.INBAND,
        antenna_pol=AntennaPol.H,
        dut_axis=DUTAxis.Z,
        dut_ant_position=DUTAntPosition.LOWER,
    )

    writer.write_inband(inband_case, InbandResult(inband_peak_level_dbm=None))

    worksheet = load_workbook(path)["LTE_BAND 1"]
    assert worksheet["A7"].value == "B1"
    assert worksheet["B7"].value == 18300
    assert worksheet["C7"].value == 20
    assert worksheet["D7"].value == 1
    assert worksheet["E7"].value == 49
    assert worksheet["F7"].value == "QPSK"
    assert worksheet["G7"].value == 15
    assert worksheet["H7"].value == "Lower"
    assert worksheet["I7"].value is None
    assert worksheet["D20"].value == "-"

    harmonic_case = _make_lte_case(
        antenna_pol=AntennaPol.H,
        dut_axis=DUTAxis.Z,
        dut_ant_position=DUTAntPosition.LOWER,
    )
    writer.write_fundamental(
        harmonic_case,
        FundamentalResult(fundamental_frequency_mhz=1950.125),
    )
    writer.write_harmonic(
        harmonic_case,
        HarmonicResult(
            harmonic_2nd_peak_level_dbm=-41.0,
            harmonic_3rd_peak_level_dbm=-40.0,
            harmonic_4th_peak_level_dbm=-39.0,
            harmonic_5th_peak_level_dbm=-38.0,
        ),
    )

    worksheet = load_workbook(path, data_only=False)["LTE_BAND 1"]
    assert worksheet["I7"].value == 1950.125
    assert [worksheet[cell].value for cell in ("A24", "A26", "A28", "A30")] == [
        "=I7*2",
        "=I7*3",
        "=I7*4",
        "=I7*5",
    ]
    assert [worksheet[cell].value for cell in ("D25", "D27", "D29", "D31")] == [
        -41.0,
        -40.0,
        -39.0,
        -38.0,
    ]
    assert writer.read_fundamental_frequency(harmonic_case) == 1950.125


def test_write_lte_fundamental_overwrites_band_and_scs(tmp_path: Path) -> None:
    path = _make_lte_workbook(tmp_path / "lte.xlsx", LTEBand.B1)
    writer = ExcelReportWriter(path)
    worksheet = writer.workbook["LTE_BAND 1"]
    worksheet["A7"] = "stale-band"
    worksheet["G7"] = 99

    writer.write_fundamental(
        _make_lte_case(),
        FundamentalResult(fundamental_frequency_mhz=1950.125),
    )

    worksheet = load_workbook(path)["LTE_BAND 1"]
    assert worksheet["A7"].value == "B1"
    assert worksheet["G7"].value == 15


def test_write_lte_b34_uses_mhb_sheet_and_bilog_cells(tmp_path: Path) -> None:
    path = _make_lte_workbook(tmp_path / "lte_b34.xlsx", LTEBand.B34)
    writer = ExcelReportWriter(path)
    case = _make_lte_case(
        band=LTEBand.B34,
        bandwidth=10.0,
        test_type=MeasurementTestType.BILOG,
        rb_allocation=LTERBAllocation.LOW,
    )

    writer.write_bilog(case, BilogResult(bilog_peak_level_dbm=-38.25))

    worksheet = load_workbook(path)["LTE_BAND 34"]
    assert worksheet["B7"].value == 36275
    assert worksheet["C7"].value == 10
    assert worksheet["B14"].value == -38.25
    assert worksheet["A24"].value == "=I7*2"
    assert worksheet["I7"].value is None


def test_write_lte_lb_uses_shifted_rows_and_preserves_formulas(tmp_path: Path) -> None:
    path = _make_lte_workbook(tmp_path / "lte_b5.xlsx", LTEBand.B5)
    writer = ExcelReportWriter(path)
    inband_case = _make_lte_case(
        band=LTEBand.B5,
        bandwidth=10.0,
        test_type=MeasurementTestType.INBAND,
        antenna_pol=AntennaPol.H,
        dut_axis=DUTAxis.Z,
        dut_ant_position=DUTAntPosition.LOWER,
        rb_allocation=LTERBAllocation.HIGH,
    )
    harmonic_case = replace(inband_case, test_type=MeasurementTestType.HARMONIC)

    writer.write_inband(inband_case, InbandResult(inband_peak_level_dbm=-47.0))
    writer.write_fundamental(
        harmonic_case,
        FundamentalResult(fundamental_frequency_mhz=836.625),
    )
    writer.write_harmonic(
        harmonic_case,
        HarmonicResult(
            harmonic_2nd_peak_level_dbm=-42.0,
            harmonic_3rd_peak_level_dbm=None,
            harmonic_4th_peak_level_dbm=-40.0,
            harmonic_5th_peak_level_dbm=-39.0,
        ),
    )

    worksheet = load_workbook(path, data_only=False)["LTE_BAND 5"]
    assert worksheet["B7"].value == 20525
    assert worksheet["E7"].value == 49
    assert worksheet["D15"].value == -47.0
    assert worksheet["I7"].value == 836.625
    assert [worksheet[cell].value for cell in ("A19", "A21", "A23", "A25")] == [
        "=I7*2",
        "=I7*3",
        "=I7*4",
        "=I7*5",
    ]
    assert [worksheet[cell].value for cell in ("D20", "D22", "D24", "D26")] == [
        -42.0,
        "-",
        -40.0,
        -39.0,
    ]


def test_prepare_live_monitoring_accepts_lte_sheet_with_extra_wcdma_sheet(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = _make_lte_workbook(
        tmp_path / "lte.xlsx",
        LTEBand.B1,
        extra_sheets=("WCDMA_BAND 1",),
    )
    writer = ExcelReportWriter(path)
    _install_fake_com_modules(monkeypatch)
    excel = _FakeExcel([])
    monkeypatch.setattr(excel_writer, "_get_excel_application", lambda *_: excel)

    writer.prepare_live_monitoring(_make_lte_case())

    opened_workbook = excel.Workbooks[-1]
    assert opened_workbook.Worksheets("LTE_BAND 1").activated
    assert writer._live_monitoring_enabled


def test_prepare_live_monitoring_rejects_nr_workbook_for_lte_case(
    tmp_path: Path,
) -> None:
    path = _make_workbook(tmp_path / "nr.xlsx", NRBand.N1)
    writer = ExcelReportWriter(path)

    with pytest.raises(ValueError, match="Workbook has no sheet"):
        writer.prepare_live_monitoring(_make_lte_case())


def test_prepare_live_monitoring_rejects_lte_workbook_for_nr_case(
    tmp_path: Path,
) -> None:
    path = _make_lte_workbook(tmp_path / "lte.xlsx", LTEBand.B1)
    writer = ExcelReportWriter(path)

    with pytest.raises(ValueError, match="Workbook has no sheet"):
        writer.prepare_live_monitoring(
            _make_case(
                band=NRBand.N1,
                bandwidth=20.0,
                test_type=MeasurementTestType.INBAND,
            )
        )
