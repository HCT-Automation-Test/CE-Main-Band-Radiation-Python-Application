from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.models.enums import Generation
from app.report import rawdata_workbook
from app.report.rawdata_source import RawDataWorkbookKind
from app.report.rawdata_workbook import (
    RawDataWorkbookError,
    ensure_rawdata_result_workbook,
    validate_eut_name,
)

_RESOURCE_ROOT = Path(__file__).resolve().parents[2]
_NR_SOURCE = _RESOURCE_ROOT / "resources" / "templates" / "NR_Rawdata_.xlsx"
_LTE_SOURCE = (
    _RESOURCE_ROOT
    / "resources"
    / "templates"
    / "_CE MAIN Radiation Auto Test Rawdata(3G,4G).xlsx"
)


@pytest.mark.parametrize(
    ("generation", "suffix", "source_path", "kind"),
    [
        (Generation.NR_SA, "NR.xlsx", _NR_SOURCE, RawDataWorkbookKind.NR),
        (Generation.LTE, "LTE.xlsx", _LTE_SOURCE, RawDataWorkbookKind.LTE),
    ],
)
def test_new_result_is_copied_to_canonical_path(
    tmp_path: Path,
    generation: Generation,
    suffix: str,
    source_path: Path,
    kind: RawDataWorkbookKind,
) -> None:
    source_hash = _sha256(source_path)

    result = ensure_rawdata_result_workbook(
        screenshot_root_dir=tmp_path,
        eut_name="EUT-A",
        generation=generation,
        resource_root=_RESOURCE_ROOT,
    )

    assert result.kind is kind
    assert result.created
    assert result.path == (
        tmp_path / "EUT-A" / f"EUT-A_CE Main_Radiation_{suffix}"
    ).resolve()
    assert _sha256(result.path) == source_hash
    assert _sha256(source_path) == source_hash


def test_nr_sa_and_nsa_share_the_same_result(tmp_path: Path) -> None:
    first = ensure_rawdata_result_workbook(
        screenshot_root_dir=tmp_path,
        eut_name="Shared-EUT",
        generation=Generation.NR_SA,
        resource_root=_RESOURCE_ROOT,
    )
    second = ensure_rawdata_result_workbook(
        screenshot_root_dir=tmp_path,
        eut_name="Shared-EUT",
        generation=Generation.NR_NSA,
        resource_root=_RESOURCE_ROOT,
    )

    assert first.created
    assert not second.created
    assert second.path == first.path


def test_existing_result_is_reused_without_overwrite(tmp_path: Path) -> None:
    first = ensure_rawdata_result_workbook(
        screenshot_root_dir=tmp_path,
        eut_name="Reuse-EUT",
        generation=Generation.NR_SA,
        resource_root=_RESOURCE_ROOT,
    )
    workbook = load_workbook(first.path)
    workbook[workbook.sheetnames[0]]["A1"] = "existing measurement"
    workbook.save(first.path)
    workbook.close()
    changed_hash = _sha256(first.path)

    second = ensure_rawdata_result_workbook(
        screenshot_root_dir=tmp_path,
        eut_name="Reuse-EUT",
        generation=Generation.NR_NSA,
        resource_root=_RESOURCE_ROOT,
    )

    assert not second.created
    assert _sha256(second.path) == changed_hash
    reused = load_workbook(second.path, read_only=True)
    try:
        assert reused[reused.sheetnames[0]]["A1"].value == "existing measurement"
    finally:
        reused.close()


def test_readable_result_is_reused_while_file_is_open(tmp_path: Path) -> None:
    first = ensure_rawdata_result_workbook(
        screenshot_root_dir=tmp_path,
        eut_name="Open-EUT",
        generation=Generation.LTE,
        resource_root=_RESOURCE_ROOT,
    )

    with first.path.open("rb"):
        second = ensure_rawdata_result_workbook(
            screenshot_root_dir=tmp_path,
            eut_name="Open-EUT",
            generation=Generation.LTE,
            resource_root=_RESOURCE_ROOT,
        )

    assert not second.created
    assert second.path == first.path


@pytest.mark.parametrize(
    "eut_name",
    [
        "",
        "   ",
        " leading",
        "trailing ",
        "trailing.",
        ".",
        "..",
        "folder/name",
        "bad:name",
        "bad*name",
        "CON",
        "con.txt",
        "LPT9",
    ],
)
def test_unsafe_eut_name_is_rejected(eut_name: str) -> None:
    with pytest.raises(RawDataWorkbookError):
        validate_eut_name(eut_name)


def test_screenshot_root_file_is_rejected(tmp_path: Path) -> None:
    root = tmp_path / "screenshots"
    root.write_text("not a directory", encoding="utf-8")

    with pytest.raises(RawDataWorkbookError, match="not a directory"):
        ensure_rawdata_result_workbook(
            screenshot_root_dir=root,
            eut_name="EUT",
            generation=Generation.NR_SA,
            resource_root=_RESOURCE_ROOT,
        )


def test_result_directory_file_is_rejected(tmp_path: Path) -> None:
    (tmp_path / "EUT").write_text("not a directory", encoding="utf-8")

    with pytest.raises(RawDataWorkbookError, match="could not be created"):
        ensure_rawdata_result_workbook(
            screenshot_root_dir=tmp_path,
            eut_name="EUT",
            generation=Generation.NR_SA,
            resource_root=_RESOURCE_ROOT,
        )


@pytest.mark.parametrize("contents", [b"broken", b""])
def test_corrupt_existing_result_is_rejected(
    tmp_path: Path,
    contents: bytes,
) -> None:
    result_dir = tmp_path / "EUT"
    result_dir.mkdir()
    result_path = result_dir / "EUT_CE Main_Radiation_NR.xlsx"
    result_path.write_bytes(contents)

    with pytest.raises(RawDataWorkbookError, match="result workbook is invalid"):
        ensure_rawdata_result_workbook(
            screenshot_root_dir=tmp_path,
            eut_name="EUT",
            generation=Generation.NR_SA,
            resource_root=_RESOURCE_ROOT,
        )


def test_wrong_managed_sheet_existing_result_is_rejected(tmp_path: Path) -> None:
    result_dir = tmp_path / "EUT"
    result_dir.mkdir()
    result_path = result_dir / "EUT_CE Main_Radiation_NR.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sub6 n999"
    workbook.save(result_path)
    workbook.close()

    with pytest.raises(RawDataWorkbookError, match="invalid NR sheet inventory"):
        ensure_rawdata_result_workbook(
            screenshot_root_dir=tmp_path,
            eut_name="EUT",
            generation=Generation.NR_SA,
            resource_root=_RESOURCE_ROOT,
        )


def test_copy_failure_removes_only_partial_result(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    result_path = (
        tmp_path / "EUT" / "EUT_CE Main_Radiation_NR.xlsx"
    ).resolve()

    def fail_copy(source_file: object, result_file: object) -> None:
        _ = source_file
        result_file.write(b"partial")  # type: ignore[attr-defined]
        raise OSError("disk full")

    monkeypatch.setattr(rawdata_workbook.shutil, "copyfileobj", fail_copy)

    with pytest.raises(RawDataWorkbookError, match="source copy failed"):
        ensure_rawdata_result_workbook(
            screenshot_root_dir=tmp_path,
            eut_name="EUT",
            generation=Generation.NR_SA,
            resource_root=_RESOURCE_ROOT,
        )

    assert not result_path.exists()
    assert _NR_SOURCE.is_file()


def test_missing_source_does_not_create_eut_directory(tmp_path: Path) -> None:
    empty_resource_root = tmp_path / "empty-resources"

    with pytest.raises(RawDataWorkbookError, match="does not exist"):
        ensure_rawdata_result_workbook(
            screenshot_root_dir=tmp_path / "screenshots",
            eut_name="EUT",
            generation=Generation.LTE,
            resource_root=empty_resource_root,
        )

    assert not (tmp_path / "screenshots" / "EUT").exists()


def test_existing_result_path_directory_is_rejected(tmp_path: Path) -> None:
    result_path = tmp_path / "EUT" / "EUT_CE Main_Radiation_NR.xlsx"
    result_path.mkdir(parents=True)

    with pytest.raises(RawDataWorkbookError, match="not a file"):
        ensure_rawdata_result_workbook(
            screenshot_root_dir=tmp_path,
            eut_name="EUT",
            generation=Generation.NR_SA,
            resource_root=_RESOURCE_ROOT,
        )


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
