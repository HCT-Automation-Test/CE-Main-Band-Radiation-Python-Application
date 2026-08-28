"""NR/LTE RawData 기준 양식 조회와 구조 검증."""
from __future__ import annotations

import hashlib
import sys
from dataclasses import dataclass
from enum import Enum
from pathlib import Path

from openpyxl import load_workbook

from app.models.enums import Generation, LTEBand, NRBand
from app.report.template_map import sheet_name_for_band

_TEMPLATE_DIRECTORY = Path("resources") / "templates"


class RawDataWorkbookKind(str, Enum):
    """RawData 결과 양식 종류."""

    NR = "NR"
    LTE = "LTE"


_SOURCE_FILENAME_BY_KIND: dict[RawDataWorkbookKind, str] = {
    RawDataWorkbookKind.NR: "NR_Rawdata_.xlsx",
    RawDataWorkbookKind.LTE: "_CE MAIN Radiation Auto Test Rawdata(3G,4G).xlsx",
}


class RawDataSourceError(RuntimeError):
    """RawData 기준 양식을 안전하게 사용할 수 없을 때 발생한다."""


@dataclass(frozen=True)
class RawDataSource:
    """검증을 마친 읽기 전용 RawData 기준 양식 정보."""

    kind: RawDataWorkbookKind
    path: Path
    sheet_names: tuple[str, ...]
    sha256: str


def workbook_kind_for_generation(generation: Generation) -> RawDataWorkbookKind:
    """Queue Generation을 RawData 양식 종류로 변환한다."""

    if generation in {Generation.NR_SA, Generation.NR_NSA}:
        return RawDataWorkbookKind.NR
    if generation is Generation.LTE:
        return RawDataWorkbookKind.LTE
    raise ValueError(f"Unsupported Generation for RawData workbook: {generation!r}")


def runtime_resource_root() -> Path:
    """개발 또는 PyInstaller 실행 환경의 resource root를 반환한다."""

    if getattr(sys, "frozen", False):
        frozen_root = getattr(sys, "_MEIPASS", None)
        if frozen_root is None:
            raise RawDataSourceError(
                "Frozen runtime does not provide the PyInstaller resource root"
            )
        return Path(frozen_root).resolve()
    return Path(__file__).resolve().parents[3]


def resolve_rawdata_source(
    generation: Generation,
    *,
    resource_root: Path | None = None,
) -> RawDataSource:
    """Generation에 대응하는 기준 양식을 찾고 읽기 가능성과 구조를 검증한다.

    ``resource_root``는 테스트와 패키징 검증을 위한 명시적 주입값이다. 런타임에서는
    개발 저장소 root 또는 PyInstaller ``sys._MEIPASS``만 사용한다.
    """

    kind = workbook_kind_for_generation(generation)
    root = runtime_resource_root() if resource_root is None else Path(resource_root)
    source_path = (
        root / _TEMPLATE_DIRECTORY / _SOURCE_FILENAME_BY_KIND[kind]
    ).resolve()
    if not source_path.is_file():
        raise RawDataSourceError(
            f"RawData source workbook does not exist: {source_path}"
        )

    sheet_names = validate_rawdata_workbook(source_path, kind)
    return RawDataSource(
        kind=kind,
        path=source_path,
        sheet_names=sheet_names,
        sha256=_sha256(source_path),
    )


def expected_sheet_names(kind: RawDataWorkbookKind) -> frozenset[str]:
    """현재 지원 Band에 필요한 관리 대상 sheet 이름을 반환한다."""

    bands = LTEBand if kind is RawDataWorkbookKind.LTE else NRBand
    return frozenset(sheet_name_for_band(band) for band in bands)


def validate_rawdata_workbook(
    path: Path,
    kind: RawDataWorkbookKind,
) -> tuple[str, ...]:
    """Workbook을 읽어 관리 대상 NR/LTE sheet 구성을 검증한다."""

    workbook_path = Path(path)
    try:
        with workbook_path.open("rb") as workbook_file:
            workbook = load_workbook(
                workbook_file,
                read_only=True,
                data_only=False,
            )
    except Exception as exc:
        raise RawDataSourceError(
            f"RawData workbook is unreadable: {workbook_path}: {exc}"
        ) from exc

    try:
        sheet_names = tuple(workbook.sheetnames)
    finally:
        workbook.close()

    expected = expected_sheet_names(kind)
    actual = set(sheet_names)
    managed_prefix = "LTE_BAND " if kind is RawDataWorkbookKind.LTE else "Sub6 "
    missing = sorted(expected - actual)
    unexpected = sorted(
        sheet_name
        for sheet_name in actual - expected
        if sheet_name.startswith(managed_prefix)
    )
    if missing or unexpected:
        details: list[str] = []
        if missing:
            details.append(f"missing={missing}")
        if unexpected:
            details.append(f"unexpected={unexpected}")
        raise RawDataSourceError(
            f"RawData workbook has an invalid {kind.value} sheet inventory: "
            f"{workbook_path}: {', '.join(details)}"
        )

    return sheet_names


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source_file:
        for chunk in iter(lambda: source_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
