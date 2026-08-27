"""결과 Excel 리포트 작성기."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.models.band_group import find_band_group
from app.models.enums import Generation
from app.models.lte_band_group import find_lte_band_group
from app.models.measurement_result import (
    BilogResult,
    FundamentalResult,
    HarmonicResult,
    InbandResult,
    MeasurementResult,
)
from app.models.nsa_anchor import validate_anchor_band_for_generation
from app.models.test_case import TestCase
from app.report.template_map import (
    ReportLayout,
    bilog_level_cell,
    harmonic_frequency_cell,
    harmonic_level_cell,
    inband_level_cell,
    layout_for_group_key_and_generation,
    sheet_name_for_band,
    validate_single_cell,
)


class ExcelReportWriter:
    """Excel 템플릿 파일에 측정 결과를 즉시 기록한다."""

    def __init__(self, template_path: Path) -> None:
        self.template_path = Path(template_path)
        self.workbook = load_workbook(self.template_path)
        self._written_band_info: set[str] = set()
        self._pending_cells: list[tuple[str, str, Any]] = []
        self._live_monitoring_enabled = False

    def prepare_live_monitoring(self, first_case: TestCase) -> None:
        """Excel을 열고 시작 TC의 Band sheet를 표시한다."""

        sheet_name = sheet_name_for_band(first_case.band)
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(
                f"Workbook has no sheet for band {first_case.band.value}: {sheet_name}"
            )

        try:
            import pythoncom
            import win32com.client
        except ImportError as exc:
            raise RuntimeError("Excel live monitoring requires pywin32") from exc

        pythoncom.CoInitialize()
        try:
            workbook = _get_or_open_excel_workbook(
                self.template_path,
                pythoncom,
                win32com.client,
            )
            _activate_excel_sheet(workbook, sheet_name)
        except Exception as exc:
            raise RuntimeError(
                f"Excel report file could not be opened for live monitoring: "
                f"{self.template_path}"
            ) from exc
        finally:
            pythoncom.CoUninitialize()

        self._live_monitoring_enabled = True

    def write_fundamental(
        self,
        test_case: TestCase,
        result: FundamentalResult,
    ) -> bool:
        """기본파 및 harmonic frequency를 작성한다."""

        frequency_mhz = _require_value(
            result.fundamental_frequency_mhz,
            "fundamental_frequency_mhz",
        )
        worksheet = self._worksheet_for_case(test_case)
        layout = self._layout_for_case(test_case)
        self._write_band_info_once(worksheet, layout, test_case)

        frequency_cell = validate_single_cell(layout.fundamental_frequency_cell)
        self._set_cell(worksheet, frequency_cell, frequency_mhz)
        if test_case.generation is Generation.LTE:
            self._save()
            return True
        for order in (2, 3, 4, 5):
            self._set_cell(
                worksheet,
                harmonic_frequency_cell(layout, order),
                frequency_mhz * order,
            )

        self._save()
        return True

    def read_fundamental_frequency(self, test_case: TestCase) -> float | None:
        """현재 Band sheet에 기록된 fundamental frequency를 읽는다."""

        worksheet = self._worksheet_for_case(test_case)
        layout = self._layout_for_case(test_case)
        frequency_cell = validate_single_cell(layout.fundamental_frequency_cell)
        value = worksheet[frequency_cell].value
        if value is None:
            return None
        if isinstance(value, bool) or not isinstance(value, int | float):
            raise ValueError(
                f"Invalid fundamental frequency value at "
                f"{worksheet.title}!{frequency_cell}: {value!r}"
            )
        return float(value)

    def write_harmonic(
        self,
        test_case: TestCase,
        result: HarmonicResult,
    ) -> bool:
        """Harmonic level 결과를 작성한다."""

        worksheet = self._worksheet_for_case(test_case)
        layout = self._layout_for_case(test_case)
        self._write_band_info_once(worksheet, layout, test_case)

        levels_by_order = {
            2: _harmonic_level_value(result.harmonic_2nd_peak_level_dbm),
            3: _harmonic_level_value(result.harmonic_3rd_peak_level_dbm),
            4: _harmonic_level_value(result.harmonic_4th_peak_level_dbm),
            5: _harmonic_level_value(result.harmonic_5th_peak_level_dbm),
        }

        for order, level_dbm in levels_by_order.items():
            self._set_cell(
                worksheet,
                harmonic_level_cell(
                    layout,
                    order,
                    test_case.dut_axis,
                    test_case.antenna_pol,
                ),
                level_dbm,
            )

        self._save()
        return True

    def write_inband(
        self,
        test_case: TestCase,
        result: InbandResult,
    ) -> bool:
        """In-band peak level 결과를 작성한다."""

        level_dbm = _harmonic_level_value(result.inband_peak_level_dbm)
        worksheet = self._worksheet_for_case(test_case)
        layout = self._layout_for_case(test_case)
        self._write_band_info_once(worksheet, layout, test_case)
        self._set_cell(
            worksheet,
            inband_level_cell(layout, test_case.dut_axis, test_case.antenna_pol),
            level_dbm,
        )
        self._save()
        return True

    def write_bilog(
        self,
        test_case: TestCase,
        result: BilogResult,
    ) -> bool:
        """Bilog peak level 결과를 작성한다."""

        level_dbm = _harmonic_level_value(result.bilog_peak_level_dbm)
        worksheet = self._worksheet_for_case(test_case)
        layout = self._layout_for_case(test_case)
        self._write_band_info_once(worksheet, layout, test_case)
        self._set_cell(
            worksheet,
            bilog_level_cell(layout, test_case.dut_axis, test_case.antenna_pol),
            level_dbm,
        )
        self._save()
        return True

    def write(
        self,
        results: list[MeasurementResult],
        output_path: Path,
    ) -> Path:
        """이전 batch API는 더 이상 지원하지 않는다."""

        raise NotImplementedError(
            "Use write_fundamental/write_harmonic/write_inband/write_bilog"
        )

    def _worksheet_for_case(self, test_case: TestCase) -> Worksheet:
        sheet_name = sheet_name_for_band(test_case.band)
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(
                f"Workbook has no sheet for band {test_case.band.value}: {sheet_name}"
            )
        return self.workbook[sheet_name]

    def _layout_for_case(self, test_case: TestCase) -> ReportLayout:
        if test_case.generation is Generation.LTE:
            group = find_lte_band_group(test_case.band)
        else:
            group = find_band_group(test_case.band)
        return layout_for_group_key_and_generation(group.key, test_case.generation)

    def _write_band_info_once(
        self,
        worksheet: Worksheet,
        layout: ReportLayout,
        test_case: TestCase,
    ) -> None:
        band_info_key = f"{worksheet.title}:{layout.key}"
        if band_info_key in self._written_band_info:
            return

        cells = layout.band_info_cells
        anchor_cell: str | None = None
        anchor_value: str | None = None
        if test_case.generation is Generation.NR_NSA:
            validate_anchor_band_for_generation(
                test_case.generation,
                test_case.anchor_band,
            )
            if cells.anchor_band is None:
                raise ValueError(
                    f"NR_NSA report layout has no anchor_band cell: {layout.key}"
                )
            anchor_cell = validate_single_cell(cells.anchor_band)
            if test_case.anchor_band is None:
                raise ValueError("NR_NSA requires anchor_band")
            anchor_value = test_case.anchor_band.value

        self._set_cell(
            worksheet,
            validate_single_cell(cells.band),
            test_case.band.value,
        )
        if anchor_cell is not None:
            self._set_cell(worksheet, anchor_cell, anchor_value)
        self._set_cell(
            worksheet,
            validate_single_cell(cells.channel),
            test_case.channel_number,
        )
        self._set_cell(
            worksheet,
            validate_single_cell(cells.bandwidth),
            test_case.bandwidth,
        )
        self._set_cell(worksheet, validate_single_cell(cells.rb_size), test_case.rb_size)
        self._set_cell(
            worksheet,
            validate_single_cell(cells.rb_offset),
            test_case.rb_offset,
        )
        self._set_cell(
            worksheet,
            validate_single_cell(cells.modulation),
            test_case.modulation.value,
        )
        self._set_cell(
            worksheet,
            validate_single_cell(cells.scs),
            test_case.scs_khz,
        )
        self._set_cell(
            worksheet,
            validate_single_cell(cells.dut_ant_position),
            test_case.dut_ant_position.value,
        )
        self._written_band_info.add(band_info_key)

    def _set_cell(self, worksheet: Worksheet, cell: str, value: Any) -> None:
        worksheet[cell] = value
        self._pending_cells.append((worksheet.title, cell, value))

    def _save(self) -> None:
        if self._live_monitoring_enabled:
            self._save_open_workbook_with_excel(open_if_missing=True)
            return

        try:
            self.workbook.save(self.template_path)
        except PermissionError:
            self._save_open_workbook_with_excel()
        else:
            self._pending_cells.clear()

    def _save_open_workbook_with_excel(self, *, open_if_missing: bool = False) -> None:
        if not self._pending_cells:
            return

        try:
            import pythoncom
            import win32com.client
        except ImportError as exc:
            raise PermissionError(
                f"Excel report file is open and cannot be saved: {self.template_path}"
            ) from exc

        pythoncom.CoInitialize()
        try:
            workbook = (
                _get_or_open_excel_workbook(
                    self.template_path,
                    pythoncom,
                    win32com.client,
                )
                if open_if_missing
                else _find_open_excel_workbook(
                    self.template_path,
                    pythoncom,
                    win32com.client,
                )
            )
            if workbook is None:
                raise PermissionError(
                    f"Excel report file is locked but no open Excel workbook was found: "
                    f"{self.template_path}"
                )

            try:
                active_sheet = self._pending_cells[-1][0]
                last_range = None
                _activate_excel_sheet(workbook, active_sheet)
                for sheet_name, cell, value in self._pending_cells:
                    worksheet = workbook.Worksheets(sheet_name)
                    target_range = worksheet.Range(cell)
                    target_range.Value = value
                    last_range = target_range
                if last_range is not None:
                    _select_excel_range(last_range)
                workbook.Save()
            except Exception as exc:
                raise PermissionError(
                    f"Excel report file is open but could not be updated through Excel: "
                    f"{self.template_path}"
                ) from exc
            self._pending_cells.clear()
        finally:
            pythoncom.CoUninitialize()


def _find_open_excel_workbook(
    template_path: Path,
    pythoncom_module: Any,
    win32_client: Any,
) -> Any | None:
    workbook = _find_workbook_from_running_object_table(
        template_path,
        pythoncom_module,
        win32_client,
    )
    if workbook is not None:
        return workbook
    return _find_workbook_from_active_excel(template_path, win32_client)


def _get_or_open_excel_workbook(
    template_path: Path,
    pythoncom_module: Any,
    win32_client: Any,
) -> Any:
    workbook = _find_open_excel_workbook(
        template_path,
        pythoncom_module,
        win32_client,
    )
    if workbook is not None:
        return workbook

    excel = _get_excel_application(win32_client)
    try:
        excel.Visible = True
        return excel.Workbooks.Open(str(template_path.resolve()))
    except Exception as exc:
        raise RuntimeError(f"Could not open Excel workbook: {template_path}") from exc


def _get_excel_application(win32_client: Any) -> Any:
    try:
        return win32_client.GetActiveObject("Excel.Application")
    except Exception:
        pass
    try:
        return win32_client.Dispatch("Excel.Application")
    except Exception as exc:
        raise RuntimeError("Could not start Excel application") from exc


def _activate_excel_sheet(workbook: Any, sheet_name: str) -> Any:
    try:
        excel = workbook.Application
        excel.Visible = True
    except Exception:
        pass
    try:
        workbook.Activate()
    except Exception:
        pass
    worksheet = workbook.Worksheets(sheet_name)
    worksheet.Activate()
    return worksheet


def _select_excel_range(target_range: Any) -> None:
    try:
        target_range.Select()
    except Exception:
        try:
            target_range.Activate()
        except Exception:
            pass


def _find_workbook_from_running_object_table(
    template_path: Path,
    pythoncom_module: Any,
    win32_client: Any,
) -> Any | None:
    target = _normalized_path(template_path)
    try:
        bind_context = pythoncom_module.CreateBindCtx(0)
        running_objects = pythoncom_module.GetRunningObjectTable()
        monikers = running_objects.EnumRunning()
    except Exception:
        return None

    while True:
        try:
            batch = monikers.Next(1)
        except Exception:
            return None
        if not batch:
            return None

        moniker = batch[0]
        try:
            candidate = running_objects.GetObject(moniker)
        except Exception:
            continue
        try:
            candidate = win32_client.Dispatch(candidate)
        except Exception:
            pass

        workbook = _match_workbook_by_path(candidate, target)
        if workbook is not None:
            return workbook


def _find_workbook_from_active_excel(template_path: Path, win32_client: Any) -> Any | None:
    try:
        excel = win32_client.GetActiveObject("Excel.Application")
    except Exception:
        return None

    target = _normalized_path(template_path)
    try:
        workbooks = excel.Workbooks
    except Exception:
        return None
    for workbook in workbooks:
        matched = _match_workbook_by_path(workbook, target)
        if matched is not None:
            return matched
    return None


def _match_workbook_by_path(candidate: Any, target_path: str) -> Any | None:
    try:
        full_name = candidate.FullName
    except Exception:
        return None
    if _normalized_path(Path(full_name)) == target_path:
        return candidate
    return None


def _normalized_path(path: Path) -> str:
    return str(path.resolve()).casefold()


def _require_value(value: float | None, field_name: str) -> float:
    if value is None:
        raise ValueError(f"Missing required report value: {field_name}")
    return value


def _harmonic_level_value(value: float | None) -> float | str:
    if value is None:
        return "-"
    return value
